import streamlit as st
import pandas as pd
import cv2
import pytesseract
from PIL import Image
from docx import Document
import re
import os
import numpy as np
from datetime import datetime

# --- SAYFA AYARLARI ---
st.set_page_config(page_title="Lazer CRM", layout="wide", page_icon="🏭")

# --- BAŞLANGIÇ VERİLERİ ---
DEFAULT_MALZEME = {
    "S235JR (Siyah)": {"fiyat": 0.85, "birim": "USD", "yogunluk": 7.85},
    "DKP": {"fiyat": 0.90, "birim": "USD", "yogunluk": 7.85},
    "Galvaniz": {"fiyat": 1.00, "birim": "USD", "yogunluk": 7.85},
    "Paslanmaz 304": {"fiyat": 3.50, "birim": "USD", "yogunluk": 7.9},
    "Paslanmaz 316": {"fiyat": 4.50, "birim": "USD", "yogunluk": 8.0},
    "Alüminyum": {"fiyat": 3.00, "birim": "USD", "yogunluk": 2.7},
    "ST37": {"fiyat": 0.85, "birim": "USD", "yogunluk": 7.85},
}

# --- SESSION STATE (HAFIZA) ---
if 'sepet' not in st.session_state: st.session_state.sepet = []
if 'malzeme_db' not in st.session_state: st.session_state.malzeme_db = DEFAULT_MALZEME
if 'dolar_kuru' not in st.session_state: st.session_state.dolar_kuru = 34.50
if 'lazer_dk_ucret' not in st.session_state: st.session_state.lazer_dk_ucret = 20.0
if 'aktif_musteri' not in st.session_state: st.session_state.aktif_musteri = None

# Form varsayılanları
defaults = {"x": 0.0, "y": 0.0, "sure": 0.0, "kal": 2.0, "fire": 0.0, "malz": "S235JR (Siyah)"}
for k, v in defaults.items():
    if f'form_{k}' not in st.session_state: st.session_state[f'form_{k}'] = v

# --- YARDIMCI FONKSİYONLAR ---

def musteri_listesi_getir():
    if os.path.exists("musteri_gecmisi.csv"):
        try:
            df = pd.read_csv("musteri_gecmisi.csv")
            isimler = df["Müşteri"].dropna().unique().tolist()
            isimler.sort()
            return isimler
        except: return []
    return []

def kayit_ekle(musteri, is_adi, tutar, detay):
    yeni_kayit = {
        "Tarih": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "Müşteri": musteri,
        "İş Adı": is_adi,
        "Tutar (TL)": round(tutar, 2),
        "Detay": detay
    }
    df = pd.DataFrame([yeni_kayit])
    mode = 'a' if os.path.exists("musteri_gecmisi.csv") else 'w'
    header = not os.path.exists("musteri_gecmisi.csv")
    df.to_csv("musteri_gecmisi.csv", mode=mode, header=header, index=False)

def sureyi_dakikaya_cevir(zaman_str):
    try:
        zaman_str = str(zaman_str).strip()
        parts = list(map(int, zaman_str.split(':')))
        if len(parts) == 3: return (parts[0] * 60) + parts[1] + (parts[2] / 60)
        elif len(parts) == 2: return parts[0] + (parts[1] / 60)
        return 0.0
    except: return 0.0

# --- ANALİZ MOTORU ---

def regex_taramasi(text):
    veriler = {}
    try:
        zaman_match = re.search(r'(?:Kesim|Cut|Time|Süre).*?(\d{2}:\d{2}:\d{2})', text, re.IGNORECASE | re.DOTALL)
        if zaman_match: veriler["sure"] = sureyi_dakikaya_cevir(zaman_match.group(1))
        
        x_match = re.search(r'[X]\s*[:|]?\s*(\d{3,5}[.,]\d+)', text)
        y_match = re.search(r'[Y]\s*[:|]?\s*(\d{3,5}[.,]\d+)', text)
        if x_match: veriler["x"] = float(x_match.group(1).replace(',', '.'))
        if y_match: veriler["y"] = float(y_match.group(1).replace(',', '.'))
        
        kal_match = re.search(r'3000\s*x\s*1500\s*x\s*(\d+[.,]?\d*)', text)
        if kal_match: veriler["kal"] = float(kal_match.group(1).replace(',', '.'))
        else:
            kal_alt = re.search(r'(?:Kalınlık|Thick|Sac)\s*[:]?\s*(\d+[.,]?\d*)', text, re.IGNORECASE)
            if kal_alt: veriler["kal"] = float(kal_alt.group(1).replace(',', '.'))

        tl = text.lower()
        if "dkp" in tl: veriler["malz"] = "DKP"
        elif "galvaniz" in tl: veriler["malz"] = "Galvaniz"
        elif "paslanmaz" in tl or "304" in tl: veriler["malz"] = "Paslanmaz 304"
        elif "alu" in tl: veriler["malz"] = "Alüminyum"
        elif "st37" in tl or "s235" in tl: veriler["malz"] = "S235JR (Siyah)"
        else: veriler["malz"] = "S235JR (Siyah)"

    except: pass
    return veriler

def word_oku(file):
    try:
        doc = Document(file)
        text_list = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                text_list.append(" ".join([cell.text for cell in row.cells]))
        return regex_taramasi("\n".join(text_list))
    except: return {}

def resim_oku(image):
    try:
        img_np = np.array(image)
        if len(img_np.shape) == 3: img_gray = cv2.cvtColor(img_np, cv2.COLOR_RGB2GRAY)
        else: img_gray = img_np
        _, img_thresh = cv2.threshold(img_gray, 150, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)
        text = pytesseract.image_to_string(Image.fromarray(img_thresh))
        return regex_taramasi(text)
    except: return {}

# --- ARAYÜZ ---

st.title("🏭 Lazer CRM Yönetim Paneli")

# MÜŞTERİ SEÇİMİ
col_m1, col_m2 = st.columns([3, 1])

with col_m1:
    musteriler = musteri_listesi_getir()
    secenekler = ["➕ Yeni Müşteri Ekle"] + musteriler
    secim = st.selectbox("Müşteri Seçimi", secenekler)

with col_m2:
    if secim == "➕ Yeni Müşteri Ekle":
        yeni_ad = st.text_input("Firma Adı", placeholder="Yeni Firma Adı...")
        if yeni_ad:
            st.session_state.aktif_musteri = yeni_ad
    else:
        st.session_state.aktif_musteri = secim
        st.success(f"**{st.session_state.aktif_musteri}**")

if not st.session_state.aktif_musteri:
    st.info("Lütfen işlem yapmak için yukarıdan bir müşteri seçin.")
    st.stop()

st.divider()

# SEKMELER
tab_is, tab_gecmis, tab_ayar = st.tabs(["🛒 Yeni Teklif", "🗂️ Müşteri Geçmişi", "⚙️ Ayarlar"])

# --- TAB 1: YENİ İŞ ---
with tab_is:
    col_sol, col_sag = st.columns([1, 1.3])
    
    with col_sol:
        st.markdown("#### 1. Veri Girişi")
        uploaded_file = st.file_uploader("Dosya Yükle (Word/Resim)", type=['docx', 'jpg', 'png', 'jpeg'])
        
        if uploaded_file:
            if "last_file" not in st.session_state or st.session_state.last_file != uploaded_file.name:
                vals = {}
                if uploaded_file.name.endswith('.docx'):
                    vals = word_oku(uploaded_file)
                    st.toast("Word okundu.", icon="📄")
                else:
                    vals = resim_oku(Image.open(uploaded_file))
                    st.toast("Resim okundu.", icon="📸")
                
                if vals:
                    for k in ['x', 'y', 'sure', 'kal', 'malz']:
                        if k in vals and vals[k] != 0:
                            st.session_state[f'form_{k}'] = vals[k]
                st.session_state.last_file = uploaded_file.name

        with st.form("veri_formu"):
            c1, c2 = st.columns(2)
            try: m_index = list(st.session_state.malzeme_db.keys()).index(st.session_state.form_malz)
            except: m_index = 0
            
            f_malz = c1.selectbox("Malzeme", list(st.session_state.malzeme_db.keys()), index=m_index)
            f_kal = c2.number_input("Kalınlık (mm)", value=float(st.session_state.form_kal))
            
            c3, c4 = st.columns(2)
            f_birim = c3.radio("Birim", ["mm", "cm", "m"], horizontal=True)
            # DÜZELTİLEN SATIR BURASI: value=1 olarak belirtildi
            f_adet = c4.number_input("Plaka Adeti", value=1, min_value=1)
            
            c5, c6 = st.columns(2)
            f_x = c5.number_input("X Boyutu", value=float(st.session_state.form_x))
            f_y = c6.number_input("Y Boyutu", value=float(st.session_state.form_y))
            
            c7, c8 = st.columns(2)
            f_sure = c7.number_input("Süre (dk)", value=float(st.session_state.form_sure))
            f_fire = c8.number_input("Fire (%)", value=0.0)
            
            if st.form_submit_button("Sepete Ekle ⬇️", type="primary", use_container_width=True):
                carpan = 1000 if f_birim == "m" else (10 if f_birim == "cm" else 1)
                st.session_state.sepet.append({
                    "Malzeme": f_malz, "K": f_kal, "X": f_x*carpan, "Y": f_y*carpan, 
                    "Süre": f_sure, "Adet": f_adet, "Fire": f_fire, "Birim": f_birim
                })
                st.rerun()

    with col_sag:
        st.markdown(f"#### 2. Sepet ({len(st.session_state.sepet)} Parça)")
        
        if st.session_state.sepet:
            df_sepet = pd.DataFrame(st.session_state.sepet)
            st.dataframe(df_sepet[["Malzeme", "K", "Adet", "Süre"]], use_container_width=True, height=150)
            
            if st.button("🗑️ Sepeti Temizle"):
                st.session_state.sepet = []
                st.rerun()
            
            top_tl = 0
            top_kg = 0
            
            for p in st.session_state.sepet:
                info = st.session_state.malzeme_db[p["Malzeme"]]
                hacim = p["X"] * p["Y"] * p["K"]
                kg = (hacim * info["yogunluk"]) / 1_000_000 * p["Adet"]
                
                fiyat = info["fiyat"] * st.session_state.dolar_kuru if info["birim"] == "USD" else info["fiyat"]
                fire_kat = 1 / (1 - p["Fire"]/100) if p["Fire"] < 100 else 1
                
                malz_tut = kg * fiyat * fire_kat
                lazer_tut = (p["Süre"] * p["Adet"]) * st.session_state.lazer_dk_ucret
                
                top_tl += malz_tut + lazer_tut
                top_kg += kg
            
            st.divider()
            c1, c2 = st.columns(2)
            c1.metric("Toplam Ağırlık", f"{top_kg:.2f} kg")
            c2.metric("Maliyet", f"{top_tl:.2f} TL")
            
            st.markdown("#### 💰 Satış Fiyatı")
            kc1, kc2 = st.columns(2)
            kar = kc1.number_input("Kâr (%)", 25, step=5)
            ekstra = kc2.number_input("Ekstra (TL)", 0)
            
            teklif = (top_tl * (1 + kar/100)) + ekstra
            st.success(f"### TEKLİF: {teklif:,.2f} TL")
            
            is_adi = st.text_input("İş Tanımı", placeholder="Örn: Flanş Kesimi")
            if st.button("💾 Kaydet", type="primary", use_container_width=True):
                kayit_ekle(st.session_state.aktif_musteri, is_adi or "Genel", teklif, f"{len(st.session_state.sepet)} parça")
                st.session_state.sepet = []
                st.success("Kayıt Başarılı!")
        else:
            st.info("Sepet boş.")

# --- TAB 2: GEÇMİŞ ---
with tab_gecmis:
    st.header(f"🗂️ {st.session_state.aktif_musteri} - İş Geçmişi")
    if os.path.exists("musteri_gecmisi.csv"):
        try:
            df_all = pd.read_csv("musteri_gecmisi.csv")
            df_mus = df_all[df_all["Müşteri"] == st.session_state.aktif_musteri]
            if not df_mus.empty:
                st.dataframe(df_mus, use_container_width=True)
                toplam = df_mus["Tutar (TL)"].sum()
                st.info(f"Toplam İş Hacmi: **{toplam:,.2f} TL**")
            else:
                st.warning("Bu müşteriye ait kayıt bulunamadı.")
        except: st.error("Veritabanı okunamadı.")
    else:
        st.warning("Kayıt yok.")

# --- TAB 3: AYARLAR ---
with tab_ayar:
    st.write("### Ayarlar")
    c1, c2 = st.columns(2)
    st.session_state.dolar_kuru = c1.number_input("Dolar Kuru", value=st.session_state.dolar_kuru)
    st.session_state.lazer_dk_ucret = c2.number_input("Lazer (TL/dk)", value=st.session_state.lazer_dk_ucret)
    if st.button("Kaydet"): st.toast("Ayarlar güncellendi.")
